"""
analytics_engine.py  --  Analytics Engineer (AE) module
Group 09 | CrediTrust Capital | UCI Statlog German Credit
==========================================================
"""

from __future__ import annotations

import argparse
import json
import os
from dataclasses import dataclass, asdict

import numpy as np
import pandas as pd


@dataclass
class CostConfig:
    c_fp: float = 1.0
    c_fn: float = 5.0
    lgd: float = 0.60
    low_risk_max: float = 0.15
    high_risk_min: float = 0.40

    @property
    def theoretical_threshold(self) -> float:
        return self.c_fp / (self.c_fp + self.c_fn)


def cost_matrix(cfg: CostConfig) -> pd.DataFrame:
    return pd.DataFrame(
        [[0.0, cfg.c_fp],
         [cfg.c_fn, 0.0]],
        index=["Actual GOOD (0)", "Actual DEFAULT (1)"],
        columns=["Decision APPROVE", "Decision REJECT"],
    )


def confusion_counts(y_true, y_pred) -> dict:
    y_true = np.asarray(y_true).astype(int)
    y_pred = np.asarray(y_pred).astype(int)
    tp = int(((y_true == 1) & (y_pred == 1)).sum())
    fp = int(((y_true == 0) & (y_pred == 1)).sum())
    fn = int(((y_true == 1) & (y_pred == 0)).sum())
    tn = int(((y_true == 0) & (y_pred == 0)).sum())
    return {"TP": tp, "FP": fp, "FN": fn, "TN": tn}


def cost_metrics(y_true, y_pred, cfg: CostConfig) -> dict:
    c = confusion_counts(y_true, y_pred)
    tp, fp, fn, tn = c["TP"], c["FP"], c["FN"], c["TN"]
    n = tp + fp + fn + tn
    precision = tp / (tp + fp) if (tp + fp) else 0.0
    recall = tp / (tp + fn) if (tp + fn) else 0.0
    denom = 2 * tp + cfg.c_fp * fp + cfg.c_fn * fn
    cw_f1 = 2 * tp / denom if denom else 0.0
    total_cost = cfg.c_fp * fp + cfg.c_fn * fn
    n_bad = tp + fn
    cost_approve_all = cfg.c_fn * n_bad
    cost_reject_all = cfg.c_fp * (n - n_bad)
    return {
        **c, "n": n,
        "precision_default": round(precision, 4),
        "recall_default": round(recall, 4),
        "cost_weighted_f1": round(cw_f1, 4),
        "total_cost": float(total_cost),
        "cost_per_applicant": round(total_cost / n, 4) if n else 0.0,
        "approval_rate": round((tn + fn) / n, 4) if n else 0.0,
        "cost_approve_all": float(cost_approve_all),
        "cost_reject_all": float(cost_reject_all),
        "cost_saved_vs_approve_all": float(cost_approve_all - total_cost),
    }


def threshold_sweep(y_true, p_default, cfg: CostConfig, thresholds=None) -> pd.DataFrame:
    if thresholds is None:
        thresholds = np.round(np.arange(0.05, 0.951, 0.01), 2)
    y_true = np.asarray(y_true).astype(int)
    p = np.asarray(p_default, dtype=float)
    rows = []
    for t in thresholds:
        pred = (p >= t).astype(int)
        m = cost_metrics(y_true, pred, cfg)
        rows.append({"threshold": float(t), **{k: m[k] for k in (
            "TP", "FP", "FN", "TN", "precision_default", "recall_default",
            "cost_weighted_f1", "total_cost", "cost_per_applicant",
            "approval_rate")}})
    return pd.DataFrame(rows)


def best_threshold(sweep: pd.DataFrame) -> float:
    best = sweep.sort_values(["total_cost", "recall_default"],
                             ascending=[True, False]).iloc[0]
    return float(best["threshold"])


def add_derived_metrics(df: pd.DataFrame, threshold: float,
                        cfg: CostConfig, p_col: str = "p_default") -> pd.DataFrame:
    out = df.copy()
    p = out[p_col].astype(float)
    out["decision"] = np.where(p >= threshold, "REJECT", "APPROVE")
    out["risk_band"] = pd.cut(
        p, bins=[-np.inf, cfg.low_risk_max, cfg.high_risk_min, np.inf],
        labels=["Low", "Medium", "High"], right=False).astype(str)
    out["exp_cost_if_approve"] = p * cfg.c_fn
    out["exp_cost_if_reject"] = (1 - p) * cfg.c_fp
    if {"credit_amount", "duration"}.issubset(out.columns):
        out["monthly_burden"] = out["credit_amount"] / out["duration"].replace(0, np.nan)
    if "credit_amount" in out.columns:
        out["expected_loss"] = p * cfg.lgd * out["credit_amount"]
        if "y_true" in out.columns:
            out["realised_loss"] = np.where(
                (out["decision"] == "APPROVE") & (out["y_true"] == 1),
                cfg.lgd * out["credit_amount"], 0.0)
    return out


def portfolio_summary(scored: pd.DataFrame) -> dict:
    s = {
        "applicants": int(len(scored)),
        "approval_rate": round(float((scored["decision"] == "APPROVE").mean()), 4),
        "high_risk_share": round(float((scored["risk_band"] == "High").mean()), 4),
    }
    if "expected_loss" in scored.columns:
        approved = scored[scored["decision"] == "APPROVE"]
        s["expected_loss_approved_book"] = round(float(approved["expected_loss"].sum()), 2)
        s["expected_loss_if_approve_all"] = round(float(scored["expected_loss"].sum()), 2)
    if "realised_loss" in scored.columns:
        s["realised_loss_approved_book"] = round(float(scored["realised_loss"].sum()), 2)
    return s


def export_threshold_sheet(sweep: pd.DataFrame, cfg: CostConfig, path: str,
                           note: str = "") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    F = "Arial"
    hdr_font = Font(name=F, bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1F3864")
    inp_font = Font(name=F, color="0000FF")
    inp_fill = PatternFill("solid", fgColor="FFFF99")
    base = Font(name=F)
    bold = Font(name=F, bold=True)

    wb = Workbook()
    a = wb.active
    a.title = "Assumptions"
    a["A1"] = "CrediTrust Capital - Cost Assumptions & Confusion Cost Matrix"
    a["A1"].font = Font(name=F, bold=True, size=13)
    a["A2"] = "Blue cells on yellow are inputs: edit them and every other sheet recalculates."
    a["A2"].font = Font(name=F, italic=True)

    rows = [
        ("Type I cost  (reject a GOOD applicant)", cfg.c_fp, "Unit cost; brief fixes ratio 1 : 5"),
        ("Type II cost (approve a DEFAULTER)", cfg.c_fn, "5x Type I, as required by the project brief"),
        ("Loss Given Default (LGD)", cfg.lgd, "ASSUMPTION - justify in report (not from dataset)"),
    ]
    a["A4"], a["B4"], a["C4"] = "Parameter", "Value", "Source / note"
    for col in "ABC":
        a[f"{col}4"].font, a[f"{col}4"].fill = hdr_font, hdr_fill
    for i, (lab, val, src) in enumerate(rows, start=5):
        a[f"A{i}"], a[f"B{i}"], a[f"C{i}"] = lab, val, src
        a[f"A{i}"].font, a[f"C{i}"].font = base, base
        a[f"B{i}"].font, a[f"B{i}"].fill = inp_font, inp_fill
    a["B7"].number_format = "0%"

    a["A9"] = "Theoretical optimal threshold  = c_fp / (c_fp + c_fn)"
    a["B9"] = "=B5/(B5+B6)"
    a["B9"].number_format = "0.000"
    a["A9"].font, a["B9"].font = bold, bold
    a["C9"] = "Reject when P(default) >= this value"
    a["C9"].font = base

    a["A11"] = "Confusion cost matrix"
    a["A11"].font = bold
    a["A12"], a["B12"], a["C12"] = "", "Decision: APPROVE", "Decision: REJECT"
    for col in "ABC":
        a[f"{col}12"].font, a[f"{col}12"].fill = hdr_font, hdr_fill
    a["A13"], a["B13"], a["C13"] = "Actual GOOD (0)", 0, "=B5"
    a["A14"], a["B14"], a["C14"] = "Actual DEFAULT (1)", "=B6", 0
    for r in (13, 14):
        for col in "ABC":
            a[f"{col}{r}"].font = base
    a.column_dimensions["A"].width = 52
    a.column_dimensions["B"].width = 20
    a.column_dimensions["C"].width = 52

    s = wb.create_sheet("Threshold Sweep")
    heads = ["Threshold", "TP", "FP (Type I)", "FN (Type II)", "TN",
             "Precision (default)", "Recall (default)", "Approval rate",
             "Total cost", "Cost / applicant", "Cost-weighted F1"]
    for j, h in enumerate(heads, start=1):
        c = s.cell(row=1, column=j, value=h)
        c.font, c.fill = hdr_font, hdr_fill
        c.alignment = Alignment(wrap_text=True, horizontal="center")
        s.column_dimensions[get_column_letter(j)].width = 16
    n_rows = len(sweep)
    for i, r in enumerate(sweep.itertuples(index=False), start=2):
        s.cell(row=i, column=1, value=r.threshold)
        s.cell(row=i, column=2, value=int(r.TP))
        s.cell(row=i, column=3, value=int(r.FP))
        s.cell(row=i, column=4, value=int(r.FN))
        s.cell(row=i, column=5, value=int(r.TN))
        s.cell(row=i, column=6, value=f"=IFERROR(B{i}/(B{i}+C{i}),0)")
        s.cell(row=i, column=7, value=f"=IFERROR(B{i}/(B{i}+D{i}),0)")
        s.cell(row=i, column=8, value=f"=(E{i}+D{i})/(B{i}+C{i}+D{i}+E{i})")
        s.cell(row=i, column=9, value=f"=C{i}*Assumptions!$B$5+D{i}*Assumptions!$B$6")
        s.cell(row=i, column=10, value=f"=I{i}/(B{i}+C{i}+D{i}+E{i})")
        s.cell(row=i, column=11,
               value=f"=IFERROR(2*B{i}/(2*B{i}+C{i}*Assumptions!$B$5+D{i}*Assumptions!$B$6),0)")
        for j in range(1, 12):
            s.cell(row=i, column=j).font = base
        s.cell(row=i, column=1).number_format = "0.00"
        for j in (6, 7, 8):
            s.cell(row=i, column=j).number_format = "0.0%"
        s.cell(row=i, column=10).number_format = "0.000"
        s.cell(row=i, column=11).number_format = "0.000"
    s.freeze_panes = "B2"
    last = n_rows + 1

    m = wb.create_sheet("Summary")
    m["A1"] = "Decision Threshold - Summary"
    m["A1"].font = Font(name=F, bold=True, size=13)
    ts = "'Threshold Sweep'"
    items = [
        ("Cost-optimal threshold (from sweep)",
         f"=INDEX({ts}!A2:A{last},MATCH(MIN({ts}!I2:I{last}),{ts}!I2:I{last},0))", "0.00"),
        ("Total cost at optimal threshold", f"=MIN({ts}!I2:I{last})", "#,##0"),
        ("Theoretical threshold c_fp/(c_fp+c_fn)", "=Assumptions!B9", "0.000"),
        ("Total cost at default 0.50 threshold",
         f"=INDEX({ts}!I2:I{last},MATCH(0.5,{ts}!A2:A{last},0))", "#,##0"),
        ("Cost saved vs. default 0.50 threshold", "=B6-B4", "#,##0"),
        ("Recall (default) at optimal threshold",
         f"=INDEX({ts}!G2:G{last},MATCH(B3,{ts}!A2:A{last},0))", "0.0%"),
        ("Approval rate at optimal threshold",
         f"=INDEX({ts}!H2:H{last},MATCH(B3,{ts}!A2:A{last},0))", "0.0%"),
        ("Cost-weighted F1 at optimal threshold",
         f"=INDEX({ts}!K2:K{last},MATCH(B3,{ts}!A2:A{last},0))", "0.000"),
    ]
    m["A2"], m["B2"] = "Metric", "Value"
    for col in "AB":
        m[f"{col}2"].font, m[f"{col}2"].fill = hdr_font, hdr_fill
    for i, (lab, fml, fmt) in enumerate(items, start=3):
        m[f"A{i}"], m[f"B{i}"] = lab, fml
        m[f"A{i}"].font, m[f"B{i}"].font = base, bold
        m[f"B{i}"].number_format = fmt
    if note:
        m["A13"] = note
        m["A13"].font = Font(name=F, italic=True)
    m.column_dimensions["A"].width = 46
    m.column_dimensions["B"].width = 18

    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb.save(path)


def _demo_frames(seed: int = 42):
    rng = np.random.default_rng(seed)

    def make(n):
        y = (rng.random(n) < 0.30).astype(int)
        p = np.clip(rng.normal(0.22 + 0.30 * y, 0.16), 0.01, 0.99)
        amount = rng.integers(500, 15000, n)
        dur = rng.integers(6, 48, n)
        return pd.DataFrame({"y_true": y, "p_default": p,
                             "credit_amount": amount, "duration": dur})
    return make(700), make(300)


def main():
    ap = argparse.ArgumentParser(description="Analytics Engineer pipeline (Group 09)")
    ap.add_argument("--tune-pred", help="CSV of validation / out-of-fold predictions")
    ap.add_argument("--test-pred", help="CSV of held-out TEST predictions")
    ap.add_argument("--out-dir", default="data/processed")
    ap.add_argument("--c-fp", type=float, default=1.0)
    ap.add_argument("--c-fn", type=float, default=5.0)
    ap.add_argument("--lgd", type=float, default=0.60)
    ap.add_argument("--demo", action="store_true")
    args = ap.parse_args()

    cfg = CostConfig(c_fp=args.c_fp, c_fn=args.c_fn, lgd=args.lgd)

    if args.demo:
        tune, test = _demo_frames()
        print("*** DEMO MODE: synthetic data, results are meaningless ***")
    else:
        if not (args.tune_pred and args.test_pred):
            ap.error("provide --tune-pred and --test-pred (or use --demo)")
        tune, test = pd.read_csv(args.tune_pred), pd.read_csv(args.test_pred)

    for name, d in (("tune", tune), ("test", test)):
        missing = {"y_true", "p_default"} - set(d.columns)
        if missing:
            raise ValueError(f"{name} file missing columns: {missing}")
        if not set(d["y_true"].unique()) <= {0, 1}:
            raise ValueError("y_true must be 0/1 with 1 = DEFAULT")

    os.makedirs(args.out_dir, exist_ok=True)

    sweep = threshold_sweep(tune["y_true"], tune["p_default"], cfg)
    t_star = best_threshold(sweep)
    sweep.to_csv(os.path.join(args.out_dir, "threshold_sweep.csv"), index=False)
    export_threshold_sheet(
        sweep, cfg, os.path.join(args.out_dir, "Decision_Threshold_Tuning_Sheet.xlsx"),
        note="Sweep computed on validation / out-of-fold predictions only.")

    res = {}
    for label, thr in (("optimal", t_star), ("default_0.5", 0.5)):
        pred = (test["p_default"] >= thr).astype(int)
        res[label] = {"threshold": thr, **cost_metrics(test["y_true"], pred, cfg)}

    scored = add_derived_metrics(test, t_star, cfg)
    scored.to_csv(os.path.join(args.out_dir, "scored_applicants.csv"), index=False)
    cost_matrix(cfg).to_csv(os.path.join(args.out_dir, "cost_matrix.csv"))

    summary = {
        "config": asdict(cfg),
        "theoretical_threshold": round(cfg.theoretical_threshold, 4),
        "chosen_threshold": t_star,
        "test_results": res,
        "portfolio_kpis": portfolio_summary(scored),
    }
    with open(os.path.join(args.out_dir, "ae_summary.json"), "w") as f:
        json.dump(summary, f, indent=2)

    print(f"Theoretical threshold : {cfg.theoretical_threshold:.3f}")
    print(f"Chosen (tuned) thresh.: {t_star:.2f}")
    for k, v in res.items():
        print(f"[{k:>11}] thr={v['threshold']:.2f}  cost={v['total_cost']:.0f}  "
              f"recall={v['recall_default']:.3f}  CW-F1={v['cost_weighted_f1']:.3f}  "
              f"approval={v['approval_rate']:.3f}")
    print(f"Files written to {args.out_dir}/")


if __name__ == "__main__":
    main()
